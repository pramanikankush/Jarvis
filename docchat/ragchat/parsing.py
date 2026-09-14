"""Parse uploaded files into (page, text) segments, then chunk them.

Formats: PDF (with OCR fallback for image-only scans), DOCX, PPTX, ODT, RTF,
HTML, JSON, CSV, TXT/MD/LOG, and images (PNG/JPG/WEBP/GIF/BMP/TIFF via OCR).
Optional/heavy paths (python-pptx, the OCR stack) import lazily so plain
-text uploads never pay for them and a missing extra degrades to a clear
message instead of a crash.
"""
import csv
import io
import logging
import os
import re

from docx import Document as _Docx
from docx.oxml.ns import qn
from docx.table import Table as _DocxTable
from pypdf import PdfReader

log = logging.getLogger("jarvis.parsing")

SUPPORTED_EXT = {
    ".pdf", ".docx", ".pptx", ".odt", ".rtf", ".html", ".htm", ".json",
    ".txt", ".md", ".markdown", ".csv", ".log",
    ".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff",
}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
# hard caps on text kept per file — protects chunking/embedding time on huge files
MAX_FILE_CHARS = 300_000
MAX_ROW_TEXT = 150_000
MIN_PDF_PAGE_CHARS = 20


def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def parse(name: str, data: bytes) -> list[tuple[int | None, str]]:
    """Return [(page_number_or_None, text), ...] for the file."""
    ext = os.path.splitext(name)[1].lower()
    if ext not in SUPPORTED_EXT:
        raise ValueError(f"Unsupported file type '{ext}'. Allowed: {', '.join(sorted(SUPPORTED_EXT))}")
    if ext == ".pdf":
        return _parse_pdf(data)
    if ext == ".docx":
        return _single(_parse_docx(data))
    if ext == ".pptx":
        return _single(_parse_pptx(data))
    if ext == ".odt":
        return _single(_parse_odt(data))
    if ext == ".rtf":
        return _single(_parse_rtf(decode_text(data)))
    if ext in (".html", ".htm"):
        return _single(_parse_html(decode_text(data)))
    if ext == ".json":
        return _single(_parse_json(decode_text(data)))
    if ext == ".csv":
        return _single(_parse_csv(data))
    if ext in IMAGE_EXTS:
        return _parse_image(data)
    text = decode_text(data)[:MAX_FILE_CHARS]
    return [] if not text.strip() else [(None, text)]


def _single(text: str) -> list[tuple[int | None, str]]:
    """Whole-file (non-paginated) formats: one segment, capped."""
    text = (text or "")[:MAX_FILE_CHARS]
    return [] if not text.strip() else [(None, text)]


def _parse_pdf(data: bytes) -> list[tuple[int, str]]:
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if len(text) >= MIN_PDF_PAGE_CHARS:
            pages.append((i, text[:MAX_FILE_CHARS]))
    if not pages:
        # no text layer anywhere -> OCR before giving up (scanned documents)
        pages = _pdf_ocr_fallback(data)
    return pages


def _pdf_ocr_fallback(data: bytes) -> list[tuple[int, str]]:
    """OCR a text-less PDF. With the OCR stack installed this reads scans;
    without it, the failure message tells the user exactly what to install."""
    from . import ocr

    if not ocr.available():
        raise ValueError("No readable text found in this PDF (it may be a scanned/image-only "
                         "document). To read scanned files, install OCR support: "
                         "pip install rapidocr-onnxruntime pypdfium2 pillow")
    log.info("PDF has no text layer — falling back to OCR")
    pages = []
    for i, text in ocr.ocr_pdf(data):
        if text.strip():
            pages.append((i, text[:MAX_FILE_CHARS]))
    if not pages:
        raise ValueError("No readable text found in this PDF — OCR could not recognize any text "
                         "(the scan may be blank, skewed, or too low-quality).")
    return pages


def _parse_image(data: bytes) -> list[tuple[int | None, str]]:
    """Screenshot / photo of a document: OCR it."""
    from . import ocr

    if not ocr.available():
        raise ValueError("Image files need OCR support: pip install rapidocr-onnxruntime pypdfium2 pillow")
    text = ocr.ocr_image_bytes(data)
    return [] if not text.strip() else [(None, text)]


def _parse_pptx(data: bytes) -> str:
    """Slide text (shapes + tables + speaker notes) in slide order."""
    try:
        from pptx import Presentation
    except ImportError:
        raise ValueError("PowerPoint support needs the 'python-pptx' package: pip install python-pptx") from None
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as e:
        raise ValueError(f"Could not read the PowerPoint file: {e}") from e
    parts = []
    for idx, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                if t.strip():
                    texts.append(t.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    line = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if line:
                        texts.append(line)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if not (texts or notes):
            continue
        block = f"[Slide {idx}]\n" + "\n".join(texts)
        if notes:
            block += f"\n(Speaker notes: {notes})"
        parts.append(block)
    return "\n\n".join(parts)


def _parse_odt(data: bytes) -> str:
    """OpenDocument text: unzip content.xml and extract paragraphs, headings,
    and table cells (stdlib zip + ElementTree — no odfpy dependency)."""
    import xml.etree.ElementTree as ET
    import zipfile

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            with z.open("content.xml") as f:
                tree = ET.parse(f)
    except (KeyError, ET.ParseError, zipfile.BadZipFile) as e:
        raise ValueError(f"Could not read the OpenDocument file: {e}") from e
    root = tree.getroot()
    body = root.find(".//{urn:oasis:names:tc:opendocument:xmlns:office:1.0}body")
    parts: list[str] = []

    def local(tag) -> str:
        return tag.rsplit("}", 1)[-1]

    def cell_text(cell) -> str:
        segs = []
        for el in cell.iter():
            if local(el.tag) in ("p", "h"):
                t = "".join(el.itertext()).strip()
                if t:
                    segs.append(t)
        return " | ".join(segs)

    def walk(el):
        tag = local(el.tag)
        if tag in ("p", "h"):
            t = "".join(el.itertext()).strip()
            if t:
                parts.append(t)
            return  # do not recurse: paragraphs are atomic units
        if tag == "table-cell":
            t = cell_text(el)
            if t:
                parts.append(t)
            return
        for child in el:
            walk(child)

    walk(body if body is not None else root)
    return "\n\n".join(parts)


_RTF_CONTROL = re.compile(r"\\([a-zA-Z]+)(-?\d+)? ?")


def _parse_rtf(text: str) -> str:
    """Pragmatic RTF text extraction (stdlib): skip destination groups (font
    tables, metadata, images), translate \\par/\\tab, unescape literals and
    \\'xx hex bytes. Good enough for RAG chunking of WordPad / "Save as RTF"
    exports — not a full RTF spec implementation."""
    _DESTINATIONS = {"fonttbl", "colortbl", "stylesheet", "info", "pict",
                     "object", "header", "footer", "headerl", "headerr",
                     "footerl", "footerr", "ftnsep", "ftnsepc", "generator",
                     "listtable", "listoverridetable", "rsidtbl", "xmlnstbl"}
    out: list[str] = []
    stack: list[str] = []  # destination name per open group ("" = plain group)
    i, n = 0, len(text)
    while i < n:
        ch = text[i]
        if ch == "{":
            j = i + 1
            dest = ""
            if j < n and text[j:j + 2] == "\\*":
                j += 2
            if j < n and text[j:j + 1] == "\\":
                m = re.match(r"\\([a-zA-Z]+)", text[j:])
                if m and m.group(1).lower() in _DESTINATIONS:
                    dest = m.group(1).lower()
            stack.append(dest)
            i += 1
        elif ch == "}":
            if stack:
                stack.pop()
            i += 1
        elif stack and any(stack):
            i += 1  # inside a destination group: drop everything
        elif ch == "\\":
            if text[i:i + 2] == "\\*":
                i += 2
            elif i + 1 < n and text[i + 1] in ("\\", "{", "}"):
                out.append(text[i + 1])
                i += 2
            elif i + 1 < n and text[i + 1] == "~":
                out.append(" ")
                i += 2
            elif i + 3 < n and text[i + 1] == "'":
                try:
                    out.append(bytes([int(text[i + 2:i + 4], 16)]).decode("cp1252", "replace"))
                except ValueError:
                    pass
                i += 4
            else:
                m = _RTF_CONTROL.match(text, i)
                if m:
                    word = m.group(1).lower()
                    if word in ("par", "line"):
                        out.append("\n")
                    elif word == "tab":
                        out.append("\t")
                    i = m.end()
                else:
                    i += 1
        elif ch in "\r\n":
            i += 1  # RTF line breaks are \\par; raw CR/LF is formatting noise
        else:
            out.append(ch)
            i += 1
    result = "".join(out)
    result = re.sub(r"[ \t]+\n", "\n", result)
    result = re.sub(r"\n{3,}", "\n\n", result)
    return result.strip()


def _parse_html(text: str) -> str:
    """Readable text from HTML — same extractor as summarize_url."""
    from .tools import extract_url_text

    return extract_url_text(text, max_chars=MAX_FILE_CHARS)


def _parse_json(text: str) -> str:
    """Flatten JSON into `path = value` lines so every value stays retrievable
    by keyword and vector search (nested objects and arrays included)."""
    import json as _json

    try:
        obj = _json.loads(text)
    except ValueError as e:
        raise ValueError(f"Invalid JSON file: {e}") from e
    lines: list[str] = []

    def render(v) -> str:
        return v if isinstance(v, str) else _json.dumps(v, ensure_ascii=False)

    def walk(node, path):
        if isinstance(node, dict):
            for k, v in node.items():
                walk(v, f"{path}.{k}" if path else str(k))
        elif isinstance(node, list):
            for idx, v in enumerate(node):
                walk(v, f"{path}[{idx}]")
        else:
            lines.append(f"{path} = {render(node)}" if path else render(node))

    walk(obj, "")
    return "\n".join(lines[:20_000])


def _parse_docx(data: bytes) -> str:
    doc = _Docx(io.BytesIO(data))
    parts = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = "".join(node.text or "" for node in child.iter(qn("w:t"))).strip()
            if text:
                parts.append(text)
        elif child.tag == qn("w:tbl"):
            table = _DocxTable(child, doc)
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                line = " | ".join(c for c in cells if c)
                if line:
                    parts.append(line)
    return "\n\n".join(parts)


def _parse_csv(data: bytes) -> str:
    text = decode_text(data)
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return ""
    header = rows[0]
    # treat first row as header only if it doesn't look like data (mostly empty or all numbers)
    nonempty = [c for c in header if c.strip()]
    looks_like_header = bool(nonempty) and not all(_is_numeric(c) for c in nonempty) and any(
        any(ch.isalpha() for ch in c) for c in nonempty
    )
    lines, total = [], 0
    for row in rows[1:] if looks_like_header else rows:
        cells = [c.strip() for c in row]
        if not any(cells):
            continue
        if looks_like_header:
            line = " | ".join(
                f"{header[i]}: {cell}" for i, cell in enumerate(cells) if i < len(header) and cell
            )
        else:
            line = ", ".join(cells)
        if not line:
            continue
        total += len(line)
        if total > MAX_ROW_TEXT:
            lines.append("[rows truncated: file is very large]")
            break
        lines.append(line)
    return "\n".join(lines)


def doc_stats(name: str) -> dict:
    """A small, honest descriptor for UI chips: what the parser accepts for
    this file type and how it is split. Page counts for PDFs come from the
    stored chunks; everything else reports its splitting mode."""
    ext = os.path.splitext(name)[1].lower()
    if ext == ".pdf":
        return {"kind": "pdf", "split_by": "page"}
    if ext == ".docx":
        return {"kind": "word", "split_by": "document"}
    if ext == ".pptx":
        return {"kind": "powerpoint", "split_by": "slide"}
    if ext == ".odt":
        return {"kind": "opendocument", "split_by": "document"}
    if ext == ".rtf":
        return {"kind": "rtf", "split_by": "document"}
    if ext in (".html", ".htm"):
        return {"kind": "html", "split_by": "document"}
    if ext == ".json":
        return {"kind": "json", "split_by": "record"}
    if ext == ".csv":
        return {"kind": "csv", "split_by": "rows"}
    if ext in IMAGE_EXTS:
        return {"kind": "image", "split_by": "ocr"}
    if ext in (".md", ".markdown"):
        return {"kind": "markdown", "split_by": "document"}
    return {"kind": "text", "split_by": "document"}


def _is_numeric(s: str) -> bool:
    try:
        float(s.replace(",", ""))
        return True
    except ValueError:
        return False


def chunk_text(text: str, max_chars: int = 1200, overlap: int = 150) -> list[str]:
    """Split text into overlapping chunks: paragraphs first, then sentences, then words."""
    if not text:
        return []
    out = []
    for para in re.split(r"\n\s*\n", text):
        para = re.sub(r"\s+", " ", para).strip()
        if not para:
            continue
        if len(para) <= max_chars:
            out.append(para)
            continue
        sentences = re.split(r"(?<=[.!?])\s+", para)
        cur = ""
        for s in sentences:
            s = s.strip()
            if not s:
                continue
            if len(s) > max_chars:
                if cur:
                    out.append(cur)
                    cur = ""
                out.extend(_hard_split(s, max_chars, overlap))
                continue
            if cur and len(cur) + len(s) + 1 > max_chars:
                out.append(cur)
                cur = cur[-overlap:] if overlap else ""
            cur = f"{cur} {s}".strip()
        if cur:
            out.append(cur)
    return out


def _hard_split(s: str, max_chars: int, overlap: int) -> list[str]:
    chunks, cur = [], ""
    for w in s.split(" "):
        if cur and len(cur) + len(w) + 1 > max_chars:
            chunks.append(cur)
            cur = cur[-overlap:] if overlap else ""
        cur = f"{cur} {w}".strip()
    if cur:
        chunks.append(cur)
    return chunks