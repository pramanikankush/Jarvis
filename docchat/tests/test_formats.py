"""Tests for the extended document parser: PPTX, ODT, RTF, HTML, JSON, logs,
OCR fallback for scanned PDFs, and image OCR (all offline; OCR-engine tests
are skipped when the optional stack is not installed).
Run: python tests/test_formats.py
"""
import io
import os
import sys
import tempfile
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from ragchat import ocr, parsing


class _PatchOcr:
    """Temporarily patch ragchat.ocr functions (parsing looks them up lazily)."""

    def __init__(self, available=True, pdf_pages=None, image_text=""):
        self._available = available
        self._pdf_pages = pdf_pages or []
        self._image_text = image_text

    def __enter__(self):
        self._orig = (ocr.available, ocr.ocr_pdf, ocr.ocr_image_bytes)
        ocr.available = lambda: self._available
        ocr.ocr_pdf = lambda data: list(self._pdf_pages)
        ocr.ocr_image_bytes = lambda data: self._image_text
        return self

    def __exit__(self, *exc):
        ocr.available, ocr.ocr_pdf, ocr.ocr_image_bytes = self._orig
        return False


def test_rtf_basic_text_and_controls():
    rtf = rb"{\rtf1\ansi Hello \par World \tab indented\par}"
    text = parsing._parse_rtf(parsing.decode_text(rtf))
    assert "Hello" in text and "World" in text, text
    assert "\t" in text and "\n" in text, text


def test_rtf_drops_destination_groups_and_escapes():
    rtf = rb"{\rtf1{\fonttbl{\f0 Arial;}}{\info{\title secret}}Hi \{brace\} Caf\'e9}"
    text = parsing._parse_rtf(parsing.decode_text(rtf))
    assert "Hi" in text and "brace" in text and "Café" in text, text
    assert "Arial" not in text and "secret" not in text, text


def test_html_strips_scripts_and_styles():
    html = ("<html><head><style>.x{color:red}</style></head><body>"
            "<h1>Title Here</h1><p>Body text</p><script>evil()</script></body></html>")
    text = parsing._parse_html(html)
    assert "Title Here" in text and "Body text" in text, text
    assert "evil()" not in text and "color:red" not in text, text


def test_json_flattens_to_retrievable_lines():
    text = parsing._parse_json('{"user": {"name": "Ada", "admin": true}, "scores": [10, 20]}')
    assert "user.name = Ada" in text, text
    assert "user.admin = true" in text, text
    assert "scores[0] = 10" in text and "scores[1] = 20" in text, text


def test_json_invalid_raises_valueerror():
    try:
        parsing._parse_json("{not json")
        assert False, "should have raised"
    except ValueError as e:
        assert "Invalid JSON" in str(e)


def test_log_and_txt_pass_through():
    pages = parsing.parse("server.log", b"2026-09-14 ERROR service down\nretrying")
    assert len(pages) == 1 and "ERROR service down" in pages[0][1]


def test_pptx_slides_tables_and_notes():
    try:
        from pptx import Presentation
    except ImportError:
        print("  skip python-pptx not installed")
        return
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarterly Review"
    slide.placeholders[1].text = "Revenue grew 42%"
    slide.notes_slide.notes_text_frame.text = "mention the Q3 spike"
    buf = io.BytesIO()
    prs.save(buf)
    pages = parsing.parse("deck.pptx", buf.getvalue())
    text = pages[0][1]
    assert "[Slide 1]" in text and "Quarterly Review" in text and "Revenue grew 42%" in text, text
    assert "Q3 spike" in text, text


def test_odt_paragraphs_and_table_cells():
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0">'
        "<office:body><office:text>"
        "<text:h>Report Title</text:h><text:p>First paragraph.</text:p>"
        '<table:table><table:table-row><table:table-cell>'
        "<text:p>CellA</text:p></table:table-cell>"
        "<table:table-cell><text:p>CellB</text:p></table:table-cell>"
        "</table:table-row></table:table>"
        "</office:text></office:body></office:document-content>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("content.xml", content)
    pages = parsing.parse("doc.odt", buf.getvalue())
    text = pages[0][1]
    assert "Report Title" in text and "First paragraph." in text, text
    # each table cell is extracted as its own segment (multi-paragraph cells
    # would be pipe-joined inside one segment)
    assert "CellA" in text and "CellB" in text, text


def test_odt_corrupt_raises_valueerror():
    buf = io.BytesIO(b"not a zip file at all")
    try:
        parsing.parse("doc.odt", buf.getvalue())
        assert False, "should have raised"
    except ValueError as e:
        assert "OpenDocument" in str(e)


def test_unsupported_extension_lists_allowed():
    try:
        parsing.parse("file.xyz", b"data")
        assert False, "should have raised"
    except ValueError as e:
        assert "Unsupported" in str(e) and ".pdf" in str(e)


def _image_only_pdf_bytes() -> bytes:
    """A PDF with no text layer: render a blank image and save it as PDF."""
    from PIL import Image

    img = Image.new("RGB", (120, 60), "white")
    buf = io.BytesIO()
    img.save(buf, format="PDF")
    return buf.getvalue()


def test_scanned_pdf_uses_ocr_fallback():
    data = _image_only_pdf_bytes()
    with _PatchOcr(available=True, pdf_pages=[(1, "SCANNED WORDS FROM OCR")]):
        pages = parsing.parse("scan.pdf", data)
    assert pages == [(1, "SCANNED WORDS FROM OCR")], pages


def test_scanned_pdf_without_ocr_gives_install_hint():
    data = _image_only_pdf_bytes()
    with _PatchOcr(available=False):
        try:
            parsing.parse("scan.pdf", data)
            assert False, "should have raised"
        except ValueError as e:
            assert "pip install" in str(e) and "scanned" in str(e)


def test_ocr_failure_on_scan_is_a_clean_error():
    data = _image_only_pdf_bytes()
    with _PatchOcr(available=True, pdf_pages=[(1, "")]):
        try:
            parsing.parse("scan.pdf", data)
            assert False, "should have raised"
        except ValueError as e:
            assert "OCR could not recognize" in str(e)


def test_image_file_goes_through_ocr():
    with _PatchOcr(available=True, image_text="SCREENSHOT TEXT"):
        pages = parsing.parse("shot.png", b"\x89PNG fake bytes")
    assert pages == [(None, "SCREENSHOT TEXT")], pages


def test_image_without_ocr_gives_install_hint():
    with _PatchOcr(available=False):
        try:
            parsing.parse("shot.png", b"\x89PNG fake bytes")
            assert False, "should have raised"
        except ValueError as e:
            assert "OCR support" in str(e)


def test_ocr_engine_smoke_when_installed():
    """When the OCR stack is present, a blank image must OCR to '' without
    raising (engine loads its bundled models on first use)."""
    if not ocr.available():
        print("  skip OCR stack not installed")
        return
    import numpy as np
    from PIL import Image

    blank = np.asarray(Image.new("RGB", (64, 32), "white"))
    assert ocr.ocr_image(blank) == ""
    assert ocr.ocr_image_bytes(b"definitely not an image") == ""  # never raises


def test_ocr_result_shapes_across_rapidocr_versions():
    """RapidOCR 1.x returns (rows, elapse) with rows shaped [box, text, score];
    3.x returns one result object exposing .txts. Both must produce text —
    unpacking the 3.x object raised TypeError and OCR'd every scan to ""."""

    class _Result3x:  # rapidocr 3.x RapidOCROutput
        txts = ("Scanned invoice", "Total 4500 USD")

    class _Empty3x:
        txts = ()

    class _Engine:
        def __init__(self, out):
            self.out = out

        def __call__(self, _img):
            return self.out

    rows = [[[[0, 0], [1, 0], [1, 1], [0, 1]], "Hello scan", 0.99], [None, "second line", 0.5]]
    assert ocr._lines_from_result(ocr._call_engine(_Engine((rows, [0.1])), None)) == \
        ["Hello scan", "second line"]
    assert ocr._lines_from_result(ocr._call_engine(_Engine(_Result3x()), None)) == \
        ["Scanned invoice", "Total 4500 USD"]
    # nothing recognized / odd rows: never raises, just no text
    assert ocr._lines_from_result(ocr._call_engine(_Engine((None, [0.0])), None)) == []
    assert ocr._lines_from_result(ocr._call_engine(_Engine(_Empty3x()), None)) == []
    assert ocr._lines_from_result(None) == []


def test_ocr_engine_failure_returns_empty_not_raises():
    """Contract: OCR is a fallback that never raises. A failing engine (missing
    or corrupt models, OOM while loading) must yield "" so the caller reports
    "no readable text" instead of turning into a 500."""
    import numpy as np

    saved = ocr._get_engine

    def _boom():
        raise RuntimeError("engine unavailable")

    ocr._get_engine = _boom
    try:
        blank = np.zeros((8, 8, 3), dtype=np.uint8)
        assert ocr.ocr_image(blank) == ""
    finally:
        ocr._get_engine = saved


def test_doc_stats_new_kinds():
    assert parsing.doc_stats("d.pptx")["kind"] == "powerpoint"
    assert parsing.doc_stats("d.json")["kind"] == "json"
    assert parsing.doc_stats("d.png")["kind"] == "image"
    assert parsing.doc_stats("d.html")["kind"] == "html"
    assert parsing.doc_stats("d.odt")["kind"] == "opendocument"
    assert parsing.doc_stats("d.rtf")["kind"] == "rtf"


def test_end_to_end_ingest_of_new_format(tmp=None):
    """ingest_doc (upload/attach pipeline) accepts the new formats."""
    import asyncio

    import numpy as np

    from ragchat import ingest
    from ragchat.store import Store

    def _embed(_t):
        return np.ones(8, dtype=np.float32) / np.sqrt(8)

    with tempfile.TemporaryDirectory() as td:
        st = Store(os.path.join(td, "t.db"))
        try:
            doc, n = asyncio.run(
                ingest.ingest_doc(st, "notes.json", b'{"topic": "quarterly results"}', _embed))
            assert doc["name"] == "notes.json" and n >= 1
        finally:
            st.close()


def main():
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    failed = 0
    for t in tests:
        try:
            t()
            print(f"  ok  {t.__name__}")
        except Exception as e:
            failed += 1
            import traceback
            print(f"FAIL  {t.__name__}: {e}")
            traceback.print_exc()
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
